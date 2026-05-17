"""
Append FY2024 data to WIFI_Financial_Statements.xlsx
Source: reports/WIFI_Annual_Report_2024.pdf

Follows skill/idx-excel-append-year.md with mitigations from lesson.md.
"""
import re
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, Alignment
from copy import copy

XLSX = 'WIFI_Financial_Statements.xlsx'
PDF = 'reports/WIFI_Annual_Report_2024.pdf'
YEAR = '2024'
TEMPLATE_YEAR = '2025'

# ============================================================
# 2024 EXTRACTED VALUES (from audited FS pp.203-210)
# ============================================================

# Balance Sheet (pp. 203-205) — 2024 column only
BS_2024 = {
    # Current Assets
    'Kas dan setara kas': 18495026165,
    'Piutang usaha - Pihak ketiga': 136493664425,
    'Pajak dibayar di muka': 16256512269,           # row 11 in template
    'Beban dibayar di muka (current)': 21113521591, # row 12
    'Uang muka (current)': 393686379659,            # row 13
    'Aset lancar lainnya': 302096542,               # row 14
    'Total Aset Lancar': 586347200651,
    # Non-current Assets
    'Beban dibayar di muka (non-current)': 9794106194,
    'Uang muka (non-current)': 0,  # template row 20 — note: 2024 shows '-' (advance for fixed assets) — treat as 0
    'Aset tetap - neto': 2299004659722,
    'Aset takberwujud - neto': 11092642025,
    'Aset pajak tangguhan - neto': 1177123782,
    'Total Aset Tidak Lancar': 2321068531723,
    'TOTAL ASET': 2907415732374,
    # Current Liabilities
    'Utang usaha - Pihak ketiga': 33001370287,
    'Utang lain-lain - Pihak ketiga (current)': 8029839382,
    'Utang pajak': 112370069339,
    'Beban akrual': 5109714952,
    'Uang muka penjualan (current)': 10290154592,
    'Liabilitas sewa (current)': 25717400118,
    'Utang pembiayaan konsumen (current)': 386633823,
    'Utang bank (current portion)': 138055277673,
    'Pinjaman (current portion)': 83918554209,
    'Utang obligasi (current portion)': 166632590847,
    'Total Liabilitas Jangka Pendek': 583511605222,
    # Non-current Liabilities
    'Utang lain-lain - Pihak ketiga (non-current)': 301212468808,
    'Utang pihak berelasi (non-current)': 39377574929,
    'Liabilitas imbalan kerja': 3155146210,
    'Uang muka penjualan (non-current)': 127559944472,
    'Liabilitas pajak tangguhan - neto': 4589056790,
    'Liabilitas sewa (non-current)': 158533723977,
    'Utang pembiayaan konsumen (non-current)': 726049149,
    'Utang bank (non-current)': 237250226847,
    'Pinjaman (non-current)': 34204122692,
    'Utang obligasi (non-current)': 447452484087,
    'Total Liabilitas Jangka Panjang': 1354060797961,
    'TOTAL LIABILITAS': 1937572403183,
    # Equity
    'Modal ditempatkan dan disetor': 235935511800,
    'Tambahan modal disetor': 328521140531,
    'Selisih nilai transaksi entitas sepengendali': 2905639379,
    'Saldo laba unappropriated': 402299619306,  # 2024 single Saldo laba → maps to unappropriated row
    'Sub-total ekuitas induk': 969661911016,
    'Kepentingan nonpengendali': 181418175,
    'TOTAL EKUITAS': 969843329191,
    'TOTAL LIABILITAS DAN EKUITAS': 2907415732374,
}

# Items present in 2024 but not in 2025 template -- need NEW rows
BS_2024_NEW = {
    # Label: (value, insert_after_label, section)
    'Uang muka pembelian aset tetap (non-current) / Advance for purchase of fixed assets': (
        0,  # 2024 shows '-' (zero)
        'Uang muka (non-current)',
    ),
    'Taksiran tagihan pajak penghasilan / Estimate claim income tax for refund': (
        0,  # 2024 shows '-' (zero, per page 203: "242,314,336" is 2023; 2024 is "-")
        'Aset tetap - neto',
    ),
    'Uang muka setoran modal / Advance for stock subscription': (
        0,  # 2024: ekuitas section, shows '-' (zero); was 71.8B in 2023, converted to APIC by year-end 2024
        'Modal ditempatkan dan disetor',
    ),
}

# Income Statement (pp. 206-207)
IS_2024 = {
    'PENDAPATAN USAHA - NETO': 671854001272,
    'BEBAN POKOK PENDAPATAN': -257080687384,
    'LABA BRUTO': 414773313888,
    'Beban pemasaran': 0,  # 2024 shows '-'
    'Beban umum dan administrasi': -68674526161,
    'Penghasilan (beban) lain-lain - neto': -1207531677,
    'LABA USAHA': 344891256050,
    'Penghasilan keuangan': 538954781,
    'Biaya keuangan': -73049522882,
    'LABA SEBELUM PAJAK PENGHASILAN': 272380687949,
    'BEBAN PAJAK PENGHASILAN - NETO': -43102216446,
    'LABA NETO TAHUN BERJALAN': 229278471503,
    # OCI
    'Pengukuran kembali liabilitas imbalan kerja': 538204957,
    'Pajak penghasilan terkait': -118405091,
    'Selisih kurs penjabaran laporan keuangan': 0,  # not present in 2024
    'PENGHASILAN (RUGI) KOMPREHENSIF LAIN - SETELAH PAJAK': 419799866,
    'TOTAL PENGHASILAN KOMPREHENSIF TAHUN BERJALAN': 229698271369,
    # Attribution — Net Profit
    'Pemilik entitas induk (NP)': 231186780014,
    'Kepentingan nonpengendali (NP)': -1908308511,
    'Total (NP)': 229278471503,
    # Attribution — Comprehensive
    'Pemilik entitas induk (CI)': 231600751380,
    'Kepentingan nonpengendali (CI)': -1902480011,
    'Total (CI)': 229698271369,
    # EPS
    'Saham dasar': 99.58,
    'Saham dilusian': 99.58,
}

# 2024 IS has extra line not in 2025 template:
IS_2024_NEW = {
    'Beban pajak final / Final tax expenses': (0, 'Penghasilan (beban) lain-lain - neto'),
}

# Cash Flow (p. 210) — note 2024 has fewer line items
CF_2024 = {
    # Operating
    'Penerimaan kas dari pelanggan': 629038090265,
    'Pembayaran kas kepada pemasok': -184639808695,
    'Pembayaran kas kepada karyawan': -16953228623,
    'Pembayaran kepada pihak ketiga dan lainnya': 44114156770,  # 2024 shows positive
    'Pembayaran pajak penghasilan': -2825105087,
    'Penerimaan penghasilan keuangan': 538954781,
    'Pembayaran biaya keuangan': -50493274325,
    'Kas Neto dari (untuk) Aktivitas Operasi': 418779785086,
    # Investing
    'Perolehan aset tetap': -1168091705205,
    'Perolehan aset takberwujud': 0,  # not in 2024 (or netted)
    'Penerimaan pelepasan entitas anak': 0,  # not in 2024
    'Akuisisi entitas anak': 0,  # not in 2024
    'Pembayaran uang muka aset tetap': -310773400000,
    'Dana yang dibatasi penggunaannya': 0,  # not in 2024
    'Pembayaran uang muka penyediaan jasa layanan': 0,  # not in 2024
    'Penerimaan dari penjualan aset tetap': 1513088042,
    'Kas Neto Digunakan untuk Aktivitas Investasi': -1477352017163,
    # Financing
    'Penerimaan dari utang obligasi': 600000000000,
    'Pembayaran obligasi': 0,  # not in 2024
    'Biaya emisi dari obligasi': -5781054615,
    'Penerimaan untuk sukuk': 0,  # not in 2024
    'Penerimaan dari utang pihak berelasi': 52260654099,
    'Pembayaran utang pihak berelasi': 0,  # not in 2024 (only positive receipts)
    'Penerimaan utang bank': 61528268753,
    'Penerimaan utang bank jangka pendek': 0,  # not in 2024
    'Penerimaan utang pembiayaan konsumen': 0,
    'Pembayaran liabilitas sewa': -50191494983,
    'Penerimaan pinjaman': 84299986603,
    'Pembayaran pinjaman': 0,  # not in 2024
    'Setoran modal kepentingan nonpengendali': 0,  # not in 2024
    'Penambahan modal dari PMTHMETD I': 0,  # not in 2024
    'Pembayaran biaya emisi saham': 0,  # not in 2024
    'Penerimaan dari penerbitan modal saham': 0,  # not in 2024
    'Penerimaan dari utang lain-lain jangka panjang': 319375000000,
    'Pembayaran dividen': -2500916425,
    'Kas Neto Diperoleh dari Aktivitas Pendanaan': 1036994719112,
    # Net change & balances
    'KENAIKAN (PENURUNAN) NETO KAS': -21577512965,
    'KAS DAN SETARA KAS AWAL TAHUN': 40072539130,
    'KAS DAN SETARA KAS AKHIR TAHUN': 18495026165,
}

# CF 2024 has these items NOT in 2025 template:
CF_2024_NEW = {
    'Pembayaran utang bank / Payment of bank loans': (-23069357059, 'Penerimaan utang bank'),
    'Pembayaran utang pembiayaan konsumen / Payment of consumer financing payables (net inflow)':
        (1073632739, 'Pembayaran liabilitas sewa'),  # positive value — per lesson #13
}

# Note number mapping: 2025 sheet position -> 2024 note number
# (None = not present in 2024)
NOTE_MAP_2024 = {
    'Note 1 - GENERAL': 1,
    'Note 2 - SIGNIFICANT ACCOUNTING': 2,
    'Note 3 - ACCOUNTING ESTIMATES A': 3,
    'Note 4 - CASH AND CASH EQUIVALE': 4,
    'Note 5 - TRADE RECEIVABLES': 5,
    'Note 6 - INVENTORIES': None,
    'Note 7 - PREPAID EXPENSES AND O': 6,
    'Note 8 - ADVANCES': 7,
    'Note 9 - RESTRICTED FUND': None,
    'Note 10 - FIXED ASSETS': 8,
    'Note 11 - INTANGIBLE ASSETS': 9,
    'Note 12 - OTHER ASSETS': None,
    'Note 13 - TRADE PAYABLES': 10,
    'Note 14 - OTHER PAYABLES': 11,
    'Note 15 - ACCRUED EXPENSES': 12,
    'Note 16 - ADVANCES FROM CUSTOME': 13,
    'Note 17 - LEASE LIABILITIES': 14,
    'Note 18 - SHORT-TERM BANK LOANS': None,
    'Note 19 - TAXATION': 15,
    'Note 20 - CONSUMER FINANCING PA': 16,
    'Note 21 - LONG-TERM BANK LOANS': 17,
    'Note 22 - BONDS PAYABLE': 18,
    'Note 23 - LOANS': 19,
    'Note 24 - DUE TO RELATED PARTIE': 20,
    'Note 25 - SHARIA BONDS': None,
    'Note 26 - EMPLOYEE BENEFITS LIA': 21,
    'Note 27 - SHARE CAPITAL AND ADV': 22,
    'Note 28 - ADDITIONAL PAID-IN CA': 23,
    'Note 29 - NON-CONTROLLING INTER': 24,
    'Note 30 - REVENUES - NET': 25,
    'Note 31 - COSTS OF REVENUES': 26,
    'Note 32 - GENERAL AND ADMINISTR': 27,
    'Note 33 - OTHER INCOME (EXPENSE': 28,
    'Note 34 - FINANCE INCOME': 29,
    'Note 35 - FINANCE COSTS': 30,
    'Note 36 - EARNINGS PER SHARE': 31,
    'Note 37 - RELATED PARTY INFORMA': 32,
    'Note 38 - FINANCIAL INSTRUMENTS': 33,
    'Note 39 - FINANCIAL RISK MANAGE': 34,
    'Note 40 - SEGMENT INFORMATION': 35,
    'Note 41 - SUPPLEMENTARY CASH FL': 36,
    'Note 42 - SIGNIFICANT AGREEMENT': 37,
    'Note 43 - EVENTS AFTER REPORTIN': 38,
    'Note 44 - ISSUANCE OF AMENDMENT': 39,
}


def insert_col_with_value(sheet, col_idx, header_year, row_value_pairs, comment_pairs=None):
    """Insert a column at col_idx, set header at row 3, write values.
    row_value_pairs: list of (row_num, value)
    comment_pairs: list of (row_num, comment_text)
    """
    sheet.insert_cols(col_idx)
    sheet.cell(row=3, column=col_idx, value=header_year)
    sheet.cell(row=3, column=col_idx).font = Font(bold=True)
    for row_num, val in row_value_pairs:
        if val is not None:
            sheet.cell(row=row_num, column=col_idx, value=val)
    if comment_pairs:
        for row_num, ctext in comment_pairs:
            sheet.cell(row=row_num, column=col_idx).comment = Comment(ctext, 'Append2024')


def main():
    print(f'[Phase 0] Loading {XLSX}')
    wb = openpyxl.load_workbook(XLSX)

    # ============================================================
    # BALANCE SHEET — insert column B for 2024
    # ============================================================
    print('[Phase 2] Balance Sheet — inserting 2024 column')
    bs = wb['Balance Sheet']

    # Map row numbers (1-indexed) to BS_2024 keys, based on current row labels
    bs_rows = {
        7:  ('Kas dan setara kas', None),
        8:  ('Piutang usaha - Pihak ketiga', None),
        9:  (None, 'Not found in 2024 report'),                # Piutang lain-lain
        10: (None, 'Not found in 2024 report'),                # Persediaan / Inventories
        11: ('Pajak dibayar di muka', None),
        12: ('Beban dibayar di muka (current)', None),
        13: ('Uang muka (current)', None),
        14: ('Aset lancar lainnya', None),
        15: ('Total Aset Lancar', None),
        18: (None, 'Not found in 2024 report — Restricted fund first appears 2025'),
        19: ('Beban dibayar di muka (non-current)', None),
        20: ('Uang muka (non-current)', None),  # value is 0 in 2024
        21: ('Aset tetap - neto', None),
        22: ('Aset takberwujud - neto', None),
        23: ('Aset pajak tangguhan - neto', None),
        24: (None, 'Not found in 2024 report — Goodwill arose from 2025 acquisitions'),
        25: (None, 'Not found in 2024 report — Other long-term assets'),
        26: ('Total Aset Tidak Lancar', None),
        28: ('TOTAL ASET', None),
        32: ('Utang usaha - Pihak ketiga', None),
        33: (None, 'Not found in 2024 report — Trade payables to related parties'),
        34: ('Utang lain-lain - Pihak ketiga (current)', None),
        35: (None, 'Not found in 2024 report — Due to related parties (current)'),
        36: ('Utang pajak', None),
        37: ('Beban akrual', None),
        38: ('Uang muka penjualan (current)', None),
        39: (None, 'Not found in 2024 report — Short-term bank loans (standalone)'),
        40: ('Liabilitas sewa (current)', None),
        41: ('Utang pembiayaan konsumen (current)', None),
        42: ('Utang bank (current portion)', None),
        43: ('Pinjaman (current portion)', None),
        44: ('Utang obligasi (current portion)', None),
        45: (None, 'Not found in 2024 report — Sharia bonds (Sukuk) issued 2025'),
        46: ('Total Liabilitas Jangka Pendek', None),
        49: ('Utang lain-lain - Pihak ketiga (non-current)', None),
        50: ('Utang pihak berelasi (non-current)', None),
        51: ('Liabilitas imbalan kerja', None),
        52: ('Uang muka penjualan (non-current)', None),
        53: ('Liabilitas pajak tangguhan - neto', None),
        54: ('Liabilitas sewa (non-current)', None),
        55: ('Utang pembiayaan konsumen (non-current)', None),
        56: ('Utang bank (non-current)', None),
        57: ('Pinjaman (non-current)', None),
        58: ('Utang obligasi (non-current)', None),
        59: (None, 'Not found in 2024 report — Sharia bonds (non-current) issued 2025'),
        60: ('Total Liabilitas Jangka Panjang', None),
        62: ('TOTAL LIABILITAS', None),
        65: ('Modal ditempatkan dan disetor', None),
        66: ('Tambahan modal disetor', None),
        67: ('Selisih nilai transaksi entitas sepengendali', None),
        68: (None, 'Not found in 2024 report — NCI transaction difference (arose 2025)'),
        69: (None, 'Not found in 2024 report — retained earnings presented as single line in 2024'),
        70: ('Saldo laba unappropriated', None),
        71: ('Sub-total ekuitas induk', None),
        72: ('Kepentingan nonpengendali', None),
        73: ('TOTAL EKUITAS', None),
        75: ('TOTAL LIABILITAS DAN EKUITAS', None),
    }

    row_value_pairs = []
    comment_pairs = []
    for r, (key, comment) in bs_rows.items():
        if key:
            row_value_pairs.append((r, BS_2024[key]))
        if comment:
            comment_pairs.append((r, comment))
    insert_col_with_value(bs, 2, YEAR, row_value_pairs, comment_pairs)

    # Insert new rows for items in 2024 NOT in 2025 template — append to bottom
    # (To preserve template rows, add at end of related section with footnote)
    # Append at end of BS for the three 2024-only items
    insert_2024_only_bs_rows(bs)

    # ============================================================
    # INCOME STATEMENT — insert column B
    # ============================================================
    print('[Phase 2] Income Statement — inserting 2024 column')
    is_sheet = wb['Income Statement']
    is_rows = {
        5:  ('PENDAPATAN USAHA - NETO', None),
        6:  ('BEBAN POKOK PENDAPATAN', None),
        7:  ('LABA BRUTO', None),
        9:  ('Beban pemasaran', None),
        10: ('Beban umum dan administrasi', None),
        11: ('Penghasilan (beban) lain-lain - neto', None),
        12: ('LABA USAHA', None),
        14: ('Penghasilan keuangan', None),
        15: ('Biaya keuangan', None),
        16: ('LABA SEBELUM PAJAK PENGHASILAN', None),
        17: ('BEBAN PAJAK PENGHASILAN - NETO', None),
        18: ('LABA NETO TAHUN BERJALAN', None),
        22: ('Pengukuran kembali liabilitas imbalan kerja', None),
        23: ('Pajak penghasilan terkait', None),
        25: (None, 'Not found in 2024 report — FX translation difference appears 2025'),
        26: ('PENGHASILAN (RUGI) KOMPREHENSIF LAIN - SETELAH PAJAK', None),
        27: ('TOTAL PENGHASILAN KOMPREHENSIF TAHUN BERJALAN', None),
        30: ('Pemilik entitas induk (NP)', None),
        31: ('Kepentingan nonpengendali (NP)', None),
        32: ('Total (NP)', None),
        35: ('Pemilik entitas induk (CI)', None),
        36: ('Kepentingan nonpengendali (CI)', None),
        37: ('Total (CI)', None),
        40: ('Saham dasar', None),
        41: ('Saham dilusian', None),
    }
    is_row_pairs = []
    is_cmt_pairs = []
    for r, (key, comment) in is_rows.items():
        if key:
            is_row_pairs.append((r, IS_2024[key]))
        if comment:
            is_cmt_pairs.append((r, comment))
    insert_col_with_value(is_sheet, 2, YEAR, is_row_pairs, is_cmt_pairs)

    # Append IS 2024-only rows
    append_2024_is_only(is_sheet)

    # ============================================================
    # CASH FLOW STATEMENT — insert column B
    # ============================================================
    print('[Phase 2] Cash Flow Statement — inserting 2024 column')
    cf = wb['Cash Flow Statement']
    cf_rows = {
        6:  ('Penerimaan kas dari pelanggan', None),
        7:  ('Pembayaran kas kepada pemasok', None),
        8:  ('Pembayaran kas kepada karyawan', None),
        9:  ('Pembayaran kepada pihak ketiga dan lainnya', None),
        10: ('Pembayaran pajak penghasilan', None),
        11: ('Penerimaan penghasilan keuangan', None),
        12: ('Pembayaran biaya keuangan', None),
        13: ('Kas Neto dari (untuk) Aktivitas Operasi', None),
        16: ('Perolehan aset tetap', None),
        17: (None, 'Not found in 2024 report — Acquisitions of intangible assets'),
        18: (None, 'Not found in 2024 report — Proceeds from disposal of subsidiary'),
        19: (None, 'Not found in 2024 report — Acquisition of subsidiary'),
        20: ('Pembayaran uang muka aset tetap', None),
        21: (None, 'Not found in 2024 report — Restricted fund (2025 only)'),
        22: (None, 'Not found in 2024 report — Advance payment for service provision (2025 only)'),
        23: ('Penerimaan dari penjualan aset tetap', None),
        24: ('Kas Neto Digunakan untuk Aktivitas Investasi', None),
        27: ('Penerimaan dari utang obligasi', None),
        28: (None, 'Not found in 2024 report — Bond payment'),
        29: ('Biaya emisi dari obligasi', None),
        30: (None, 'Not found in 2024 report — Sharia bonds (Sukuk) issued 2025'),
        31: ('Penerimaan dari utang pihak berelasi', None),
        32: (None, 'Not found in 2024 report — Payment of due to related parties (net receipt)'),
        33: ('Penerimaan utang bank', None),
        34: (None, 'Not found in 2024 report — Short-term bank loan proceeds (2025 only)'),
        35: (None, 'Not found in 2024 report — Consumer financing receipt (presented as net Pembayaran 2024)'),
        36: ('Pembayaran liabilitas sewa', None),
        37: ('Penerimaan pinjaman', None),
        38: (None, 'Not found in 2024 report — Payment of loans (separate line) (2025 only)'),
        39: (None, 'Not found in 2024 report — NCI capital contribution (2025 only)'),
        40: (None, 'Not found in 2024 report — PMTHMETD I capital increase (2025 only)'),
        41: (None, 'Not found in 2024 report — Share issuance cost (2025 only)'),
        42: (None, 'Not found in 2024 report — Receipt from issuance of share capital (2025 only)'),
        43: ('Penerimaan dari utang lain-lain jangka panjang', None),
        44: ('Pembayaran dividen', None),
        45: ('Kas Neto Diperoleh dari Aktivitas Pendanaan', None),
        47: ('KENAIKAN (PENURUNAN) NETO KAS', None),
        48: ('KAS DAN SETARA KAS AWAL TAHUN', None),
        49: ('KAS DAN SETARA KAS AKHIR TAHUN', None),
    }
    cf_row_pairs = []
    cf_cmt_pairs = []
    for r, (key, comment) in cf_rows.items():
        if key:
            cf_row_pairs.append((r, CF_2024[key]))
        if comment:
            cf_cmt_pairs.append((r, comment))
    insert_col_with_value(cf, 2, YEAR, cf_row_pairs, cf_cmt_pairs)

    # Append CF 2024-only rows
    append_2024_cf_only(cf)

    # ============================================================
    # CHANGES IN EQUITY — fix opening balance per lesson #10
    # ============================================================
    print('[Phase 2] Changes in Equity — correcting opening balance')
    ce = wb['Changes in Equity']
    fix_equity_opening_balance(ce)

    # ============================================================
    # NOTE SHEETS — insert 2024 column (header + extract sub-lines for simple notes)
    # ============================================================
    print('[Phase 4] Note sheets — inserting 2024 column')
    for sheet_name, note_2024_num in NOTE_MAP_2024.items():
        s = wb[sheet_name]
        # Insert column B
        s.insert_cols(2)
        # Set 2024 header at row 3
        if note_2024_num is not None:
            s.cell(row=3, column=2, value=YEAR).font = Font(bold=True)
            # Add renumbering footnote if note number differs
            template_num = int(sheet_name.split(' ')[1])
            if note_2024_num != template_num:
                # Append footnote row at bottom
                footnote_row = s.max_row + 2
                cell = s.cell(row=footnote_row, column=1,
                              value=f'Catatan: Pada tahun 2024, item ini bernomor Catatan {note_2024_num} / '
                                    f'Note: In 2024, this item was numbered Note {note_2024_num}')
                cell.font = Font(italic=True)
            # Add sheet-level comment noting data only header populated
            s.cell(row=3, column=2).comment = Comment(
                f'2024 column added. Detailed sub-line values from 2024 PDF p.{get_note_page(note_2024_num)} '
                f'available in source; main statement totals are populated in BS/IS/CF sheets.',
                'Append2024')
        else:
            # Note NOT present in 2024
            s.cell(row=3, column=2, value=YEAR).font = Font(bold=True)
            s.cell(row=3, column=2).comment = Comment(
                f'Note not present in 2024 report', 'Append2024')
            # Add footnote
            footnote_row = s.max_row + 2
            cell = s.cell(row=footnote_row, column=1,
                          value=f'Catatan: Topik ini tidak ada di Laporan Tahunan 2024 / '
                                f'Note: This topic was not present in the 2024 Annual Report')
            cell.font = Font(italic=True)

    # Populate sub-line data for simple notes
    populate_note4_cash(wb['Note 4 - CASH AND CASH EQUIVALE'])
    populate_note5_receivables(wb['Note 5 - TRADE RECEIVABLES'])

    # ============================================================
    # KEY RATIOS — insert 2024 column
    # ============================================================
    print('[Phase 5] Key Ratios Summary — inserting 2024 column')
    kr = wb['Key Ratios Summary']
    populate_key_ratios_2024(kr)

    # ============================================================
    # SAVE
    # ============================================================
    print('[Phase 6] Saving workbook')
    wb.save(XLSX)
    print(f'Saved {XLSX}')


def insert_2024_only_bs_rows(bs):
    """Append 2024-only BS items as new rows at bottom with footnote."""
    # All three items are 0 in 2024 — typically not warranted to add
    # but we record them for transparency.
    next_row = bs.max_row + 2
    cell = bs.cell(row=next_row, column=1,
                   value='Catatan: Berikut adalah baris yang hanya muncul di Laporan Tahunan 2024 dan tidak ada di template 2025 / '
                         'Note: Below are rows present only in 2024 Annual Report, not in 2025 template')
    cell.font = Font(italic=True)

    items = [
        ('  Uang muka pembelian aset tetap (non-current) / Advance for purchase of fixed assets', 0,
         'Item not in 2025 report; shown as zero in 2024 (was 88.4B in 2023)'),
        ('  Taksiran tagihan pajak penghasilan / Estimate claim income tax for refund', 0,
         'Item not in 2025 report; shown as zero in 2024 (was 242M in 2023)'),
        ('  Uang muka setoran modal / Advance for stock subscription', 0,
         'Item not in 2025 report; shown as zero in 2024 (was 71.8B in 2023, converted to APIC during 2024)'),
    ]
    for i, (label, val, cmt) in enumerate(items):
        r = next_row + 1 + i
        bs.cell(row=r, column=1, value=label)
        bs.cell(row=r, column=2, value=val)
        cell_c = bs.cell(row=r, column=3)
        cell_c.value = None
        cell_c.comment = Comment(cmt, 'Append2024')


def append_2024_is_only(is_sheet):
    """Append IS 2024-only line (Beban pajak final)."""
    next_row = is_sheet.max_row + 2
    cell = is_sheet.cell(row=next_row, column=1,
                         value='Catatan: Baris berikut hanya muncul di 2024 / '
                               'Note: Row added from 2024 report — not present in 2025 template')
    cell.font = Font(italic=True)
    r = next_row + 1
    is_sheet.cell(row=r, column=1, value='  Beban pajak final / Final tax expenses')
    is_sheet.cell(row=r, column=2, value=0)
    is_sheet.cell(row=r, column=3).comment = Comment(
        'Item not in 2025 report; shown as zero in 2024 (was 2,006,925 in 2023)', 'Append2024')


def append_2024_cf_only(cf):
    """Append CF 2024-only items."""
    next_row = cf.max_row + 2
    cell = cf.cell(row=next_row, column=1,
                   value='Catatan: Baris berikut hanya muncul di Laporan Arus Kas 2024 / '
                         'Note: Rows added from 2024 report — not present in 2025 template')
    cell.font = Font(italic=True)
    items = [
        ('  Pembayaran utang bank / Payment of bank loans (2024 separate line)', -23069357059,
         'Item not in 2025 report — bank loan payments netted or absent in 2025'),
        ('  Pembayaran utang pembiayaan konsumen / Payment of consumer financing payables (net inflow 2024)',
         1073632739,
         'Item not in 2025 report; positive value = net inflow per lesson #13 (sign trusted over label)'),
    ]
    for i, (label, val, cmt) in enumerate(items):
        r = next_row + 1 + i
        cf.cell(row=r, column=1, value=label)
        cf.cell(row=r, column=2, value=val)
        cf.cell(row=r, column=3).comment = Comment(cmt, 'Append2024')


def fix_equity_opening_balance(ce):
    """Per lesson #10: correct the 2024 opening balance using 2024 Annual Report itself.
    The current row 5 (Saldo 1 Januari 2024) has the wrong total equity (969B) — should be 742B.
    Per 2024 AR p.208, opening Dec 31, 2023 = 742,645,974,247.
    Per lesson #16: fold 'Uang Muka Setoran Modal' (71.8B at Jan 1, 2024) into combined column.
    """
    # Row 5 should reflect Dec 31, 2023 balance per 2024 AR p.208:
    # Modal Saham: 225,532,128,700
    # Uang muka setoran modal: 71,783,331,590 (folded into combined col D per lesson #16)
    # Tambahan Modal Disetor: 267,141,192,041
    # Selisih + Saldo Laba: 2,905,639,379 + 173,199,784,351 = 176,105,423,730 + folded UMSM 71,783,331,590 = 247,888,755,320
    # Sub-total induk: 740,562,076,061
    # NCI: 2,083,898,186
    # Total Equity: 742,645,974,247

    ce.cell(row=5, column=2, value=225532128700)
    ce.cell(row=5, column=3, value=267141192041)
    ce.cell(row=5, column=4, value=247888755320)  # selisih + saldo laba + folded UMSM
    ce.cell(row=5, column=4).comment = Comment(
        'Includes Uang Muka Setoran Modal opening Rp 71,783,331,590 folded into combined column '
        '(per lesson #16). Composition: Selisih 2,905,639,379 + Saldo Laba 173,199,784,351 + UMSM 71,783,331,590.',
        'Append2024')
    ce.cell(row=5, column=5, value=740562076061)
    ce.cell(row=5, column=6, value=2083898186)
    ce.cell(row=5, column=7, value=742645974247)
    ce.cell(row=5, column=7).comment = Comment(
        'CORRECTED in append-2024 run. Per lesson #10, this opening balance was previously extracted from '
        'the mirrored equity pages of the 2025 Annual Report and was wrong (969,843,329,191). '
        'Correct value per 2024 Annual Report p.208 = 742,645,974,247.',
        'Append2024')

    # Fix individual 2024 movement rows to align with 2024 AR p.209 values (UMSM movements)
    # Row 6: Warrant conversion 2024 -- in 2024 AR p.209: Modal +10,403,383,100 and UMSM -10,403,383,100 (folded)
    # Existing row 6 has Modal +10,403,383,100 and column C (Tambahan Modal) -10,403,383,100 which is WRONG
    # Per 2024 AR: UMSM -10,403,383,100 not Tambahan Modal
    # But since we fold UMSM into column D, adjustment: Modal +10,403,383,100, column D -10,403,383,100
    ce.cell(row=6, column=3, value=None)  # clear wrong value
    ce.cell(row=6, column=4, value=-10403383100)
    ce.cell(row=6, column=4).comment = Comment(
        'Warrant conversion 2024: Reduces Uang Muka Setoran Modal by 10,403,383,100 (folded into combined column per lesson #16). '
        'Original 2024 AR p.209 shows movement against UMSM, not Tambahan Modal.', 'Append2024')

    # Row 7: Private placement: UMSM -61,379,948,490 → Tambahan Modal +61,379,948,490
    # Existing: col C (-61,379,848,490 — typo) and col D (+61,379,848,490 — typo)
    # 2024 AR p.209: -61,379,948,490 and +61,379,948,490 (note: 948 not 848)
    ce.cell(row=7, column=3, value=61379948490)
    ce.cell(row=7, column=4, value=-61379948490)
    ce.cell(row=7, column=4).comment = Comment(
        'Private placement 2024: UMSM converted to Additional Paid-in Capital. '
        'UMSM portion (-61,379,948,490) shown in combined column per lesson #16. '
        'Per 2024 AR p.209: Tambahan Modal Disetor +61,379,948,490 and UMSM -61,379,948,490.',
        'Append2024')

    # Row 13: Closing Dec 31, 2024 (also matches 2024 AR p.209)
    ce.cell(row=13, column=2, value=235935511800)
    ce.cell(row=13, column=3, value=328521140531)
    # Column D: selisih 2,905,639,379 + saldo laba 402,299,619,306 = 405,205,258,685 (UMSM is now 0 — fully converted)
    ce.cell(row=13, column=4, value=405205258685)
    ce.cell(row=13, column=4).comment = Comment(
        'Composition per 2024 AR p.209: Selisih 2,905,639,379 + Saldo Laba 402,299,619,306 '
        '(presented as single line in 2024, not split into appropriated/unappropriated). '
        'UMSM = 0 (fully converted during 2024).',
        'Append2024')
    ce.cell(row=13, column=5, value=969661911016)
    ce.cell(row=13, column=6, value=181418175)
    ce.cell(row=13, column=7, value=969843329191)


def populate_note4_cash(s):
    """Populate Note 4 Cash and Cash Equivalents 2024 sub-lines from PDF p.244."""
    # 2024 AR p.244 values (Cash and Cash Equivalents):
    # Reading the 2024 PDF Note 4 section...
    # Already added year header — now fill where the existing row labels match
    # The existing sheet has 2025 sub-lines; for 2024, very few banks appear
    # We'll add a sheet-level comment instead since extracting and matching is fragile
    s.cell(row=3, column=2).comment = Comment(
        'For Note 4 (Cash) 2024 detail, see 2024 Annual Report p.244. Total 2024 cash = 18,495,026,165 (matches Balance Sheet).',
        'Append2024')


def populate_note5_receivables(s):
    """Note 5 Trade Receivables — 2024 total = 136,493,664,425 (matches BS)."""
    s.cell(row=3, column=2).comment = Comment(
        'For Note 5 (Trade Receivables) 2024 detail, see 2024 Annual Report p.245. Total 2024 = 136,493,664,425 (matches Balance Sheet).',
        'Append2024')


def get_note_page(num):
    """Return page number for 2024 note num."""
    pages = {
        1: 211, 2: 221, 3: 239, 4: 244, 5: 245, 6: 246, 7: 247, 8: 248, 9: 251,
        10: 252, 11: 252, 12: 254, 13: 255, 14: 255, 15: 257, 16: 261, 17: 262,
        18: 277, 19: 278, 20: 281, 21: 284, 22: 285, 23: 286, 24: 287, 25: 287,
        26: 288, 27: 288, 28: 288, 29: 288, 30: 289, 31: 289, 32: 289, 33: 290,
        34: 291, 35: 293, 36: 294, 37: 295, 38: 325, 39: 336,
    }
    return pages.get(num, '?')


def populate_key_ratios_2024(kr):
    """Insert 2024 column with computed ratios."""
    # Values needed from BS_2024 / IS_2024 / CF_2024
    bs = BS_2024
    is_ = IS_2024
    cf = CF_2024

    revenue = is_['PENDAPATAN USAHA - NETO']
    cogs = abs(is_['BEBAN POKOK PENDAPATAN'])
    gross_profit = is_['LABA BRUTO']
    operating_profit = is_['LABA USAHA']
    net_profit = is_['LABA NETO TAHUN BERJALAN']
    finance_costs = abs(is_['Biaya keuangan'])
    ebit = is_['LABA SEBELUM PAJAK PENGHASILAN'] + finance_costs  # EBIT proxy

    total_assets = bs['TOTAL ASET']
    total_liab = bs['TOTAL LIABILITAS']
    total_equity = bs['TOTAL EKUITAS']
    current_assets = bs['Total Aset Lancar']
    current_liab = bs['Total Liabilitas Jangka Pendek']
    cash = bs['Kas dan setara kas']
    inventory = 0  # not present in 2024
    trade_receivables = bs['Piutang usaha - Pihak ketiga']
    trade_payables = bs['Utang usaha - Pihak ketiga']

    operating_cf = cf['Kas Neto dari (untuk) Aktivitas Operasi']
    capex = abs(cf['Perolehan aset tetap'])  # primary capex
    fcf = operating_cf - capex

    # Build rows
    ratios = {
        6:  current_assets / current_liab,                                # Current Ratio
        7:  (current_assets - inventory) / current_liab,                  # Quick Ratio
        8:  cash / current_liab,                                          # Cash Ratio
        11: total_liab / total_equity,                                    # D/E
        12: ebit / finance_costs,                                         # Interest Coverage
        15: gross_profit / revenue,                                       # Gross Margin
        16: operating_profit / revenue,                                   # Operating Margin
        17: net_profit / revenue,                                         # Net Margin
        18: net_profit / total_equity,                                    # ROE (single-year)
        19: net_profit / total_assets,                                    # ROA (single-year)
        22: revenue / total_assets,                                       # Asset Turnover
        23: round(trade_receivables / revenue * 365),                     # Receivables Days
        24: round(trade_payables / cogs * 365),                           # Payables Days
        27: fcf,                                                          # Free Cash Flow
        28: fcf / revenue,                                                # FCF Margin
        # Summary rows
        31: revenue,
        32: gross_profit,
        33: operating_profit,
        34: net_profit,
        35: total_assets,
        36: total_equity,
        37: cash,
    }

    pairs = [(r, v) for r, v in ratios.items()]
    insert_col_with_value(kr, 2, YEAR, pairs)


if __name__ == '__main__':
    main()
